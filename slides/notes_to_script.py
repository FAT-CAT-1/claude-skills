# -*- coding: utf-8 -*-
"""notes_to_script.py — pptx のノートから、紙用の台本を機械的に書き出す

    python notes_to_script.py deck.pptx [出力.md] [--変奏 5分版]

⚠**台本を書く場所はノートだけ**（＝ノートが唯一の正）。
  ここは「書き写す」のではなく「取り出す」ためのもの。人が別ファイルに台詞を書くと、
  必ずどちらか片方だけ直されて食い違う。だから書き出しは**常に機械的**に行う。

  紙の保険が要る理由: 本番でPCが落ちるとノートも同時に消える。
  PowerPoint の「ノート付き印刷」でも同じことができるので、好きな方でよい。

ノートの書式（deck.note() が生成する形）:
    [0:00-0:20]
    ※5分版 では飛ばす
    台詞…
    ⚠ 注意…
"""
import io
import re
import sys
from pptx import Presentation


def parse(text):
    time, skips, say, warn = None, [], [], []
    for ln in (text or '').split('\n'):
        t = ln.strip()
        if not t:
            continue
        m = re.fullmatch(r'\[(.+?)\]', t)
        if m and time is None:
            time = m.group(1)
            continue
        m = re.fullmatch(r'※(.+?) では飛ばす', t)
        if m:
            skips += m.group(1).split('／')
            continue
        if t.startswith('⚠'):
            warn.append(t.lstrip('⚠ ').strip())
            continue
        say.append(t)
    return time, skips, say, warn


def main():
    args = [a for a in sys.argv[1:] if not a.startswith('--')]
    variant = None
    for i, a in enumerate(sys.argv):
        if a == '--変奏' and i + 1 < len(sys.argv):
            variant = sys.argv[i + 1]
    if not args:
        print('使い方: python notes_to_script.py deck.pptx [出力.md] [--変奏 5分版]')
        return 2
    src = args[0]
    dst = args[1] if len(args) > 1 else src.rsplit('.', 1)[0] + '_台本.md'

    prs = Presentation(src)
    slides = []
    for i, s in enumerate(prs.slides, 1):
        note = s.notes_slide.notes_text_frame.text if s.has_notes_slide else ''
        head = ' ／ '.join(
            sh.text_frame.text.strip().replace('\n', ' ')
            for sh in s.shapes
            if sh.has_text_frame and sh.text_frame.text.strip())[:40]
        slides.append((i, head, parse(note)))

    o = io.StringIO()
    title = src.rsplit('.', 1)[0]
    o.write('# %s — 台本%s\n\n' % (title, ('（%s）' % variant) if variant else ''))
    o.write('> ⚠ **この文書は `%s` のノートから自動生成したもの。ここを直しても反映されない。**\n'
            '> 台詞を直すときは pptx のノートを直して、これを作り直す。\n\n' % src)

    shown = [(i, h, p) for (i, h, p) in slides
             if not (variant and variant in p[1])]
    o.write('| 時刻 | 枚 | 内容 |\n|---|---|---|\n')
    for i, head, (time, skips, say, warn) in shown:
        o.write('| %s | %d | %s |\n' % (time or '—', i, head or '（画面のみ）'))
    if variant:
        skipped = [str(i) for (i, _, p) in slides if variant in p[1]]
        if skipped:
            o.write('\n**%s で飛ばす: %s 枚目**\n' % (variant, '・'.join(skipped)))
    o.write('\n---\n\n')

    for i, head, (time, skips, say, warn) in shown:
        o.write('## %d枚目%s%s\n\n' % (i, ('  [%s]' % time) if time else '',
                                     ('  — %s' % head) if head else ''))
        if skips and not variant:
            o.write('> ※ %s では飛ばす\n\n' % '／'.join(skips))
        for line in say:
            o.write('> %s\n' % line)
        if not say:
            o.write('> （台詞なし／画面を見せる）\n')
        for w in warn:
            o.write('\n⚠ %s\n' % w)
        o.write('\n---\n\n')

    io.open(dst, 'w', encoding='utf-8').write(o.getvalue())
    print('台本を書き出した: %s（%d枚%s）'
          % (dst, len(shown), ('／%s で%d枚を除外' % (variant, len(slides) - len(shown)))
             if variant else ''))
    return 0


if __name__ == '__main__':
    sys.exit(main())
