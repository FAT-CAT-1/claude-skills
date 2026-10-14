# -*- coding: utf-8 -*-
"""verify.py — スライドを「実際に描いて」検査する（/slides スキルの心臓部）

    python verify.py deck.pptx [出力先フォルダ]

⚠**このスキルで最も重要なのがこの工程**。
  過去2回、文字数やはみ出しの数値チェックは全部緑だったのに、
  実際に描いて見たら折り返しが崩れていた（末尾の数文字だけが次行に取り残されていた）。
  **数値チェックは折り返しの崩れを捕まえない。** だから必ず絵にして目で見る。

描画の3段構え:
  1. PowerPoint 本体（Windows + Office）… 本番と同じ絵が出る。最優先。
  2. LibreOffice → PDF → PNG（PyMuPDF）… 近い絵が出る。書体差は出うる。
  3. どちらも無い … **「目視していない」と明言して異常終了する。**
     ここで黙って成功扱いにしない。検査していないものを検査したと言わないため。
"""
import os
import sys
import glob
import shutil
import subprocess


def _clean(out):
    os.makedirs(out, exist_ok=True)
    for f in glob.glob(os.path.join(out, '*.png')):
        os.remove(f)


def by_powerpoint(src, out, long_side=1600):
    """⚠書き出しの縦横比は**スライド自身から取る**。
    16:9 前提で 1600x900 に固定すると、4:3 のスライドが横に潰れて書き出され、
    本番と違う絵を見ながら検査することになる（実際に歪んだ絵が出た）。"""
    try:
        import win32com.client
    except ImportError:
        return None
    try:
        app = win32com.client.Dispatch('PowerPoint.Application')
    except Exception:
        return None
    pres = None
    try:
        pres = app.Presentations.Open(os.path.abspath(src), WithWindow=False)
        pw, ph = pres.PageSetup.SlideWidth, pres.PageSetup.SlideHeight
        width = long_side
        height = int(round(long_side * ph / pw))
        for i, slide in enumerate(pres.Slides, 1):
            slide.Export(os.path.join(out, '%02d.png' % i), 'PNG', width, height)
        print('  書き出し寸法: %dx%d（スライド比 %.2f:1）' % (width, height, pw / ph))
        return len(pres.Slides)
    except Exception as e:
        print('  PowerPoint での書き出しに失敗: %s' % e)
        return None
    finally:
        try:
            if pres:
                pres.Close()
            app.Quit()
        except Exception:
            pass


def by_libreoffice(src, out, dpi=150):
    soffice = shutil.which('soffice') or shutil.which('libreoffice')
    if not soffice:
        return None
    try:
        import fitz  # PyMuPDF
    except ImportError:
        print('  LibreOffice はあるが PyMuPDF が無い（pip install pymupdf）')
        return None
    try:
        subprocess.run([soffice, '--headless', '--convert-to', 'pdf',
                        '--outdir', out, os.path.abspath(src)],
                       check=True, capture_output=True, timeout=300)
    except Exception as e:
        print('  LibreOffice での変換に失敗: %s' % e)
        return None
    pdf = os.path.join(out, os.path.splitext(os.path.basename(src))[0] + '.pdf')
    if not os.path.exists(pdf):
        return None
    doc = fitz.open(pdf)
    for i, page in enumerate(doc, 1):
        page.get_pixmap(dpi=dpi).save(os.path.join(out, '%02d.png' % i))
    n = len(doc)
    doc.close()
    return n


# ── 構造チェック（絵にする前の足切り。これ"だけ"では完成にしない） ──
def structural(src):
    from pptx import Presentation
    from pptx.util import Pt
    prs = Presentation(src)
    W, H = prs.slide_width, prs.slide_height
    issues, empty_notes = [], []
    for i, s in enumerate(prs.slides, 1):
        has_text = False
        for sh in s.shapes:
            r = None
            try:
                r = (sh.left, sh.top, sh.width, sh.height)
            except Exception:
                pass
            if r and None not in r:
                if r[0] < -1000 or r[1] < -1000 or r[0] + r[2] > W + 1000 or r[1] + r[3] > H + 1000:
                    issues.append('%d枚目: 図形がスライドの外にはみ出している' % i)
            if sh.has_text_frame and sh.text_frame.text.strip():
                has_text = True
                for p in sh.text_frame.paragraphs:
                    for run in p.runs:
                        if run.font.size and run.font.size < Pt(12):
                            issues.append('%d枚目: %dpt の文字がある（後ろの席から読めない）'
                                          % (i, run.font.size.pt))
        note = s.notes_slide.notes_text_frame.text.strip() if s.has_notes_slide else ''
        if not note:
            empty_notes.append(i)
        if not has_text and not note:
            issues.append('%d枚目: 本文もノートも空（意図した空白ページか確認）' % i)
    return prs, issues, empty_notes


def contrast(src_theme_fg, src_theme_bg):
    """WCAG の相対輝度比。本文は 4.5 以上、大きな見出しは 3.0 以上が目安。"""
    def lum(h):
        v = []
        for k in (0, 2, 4):
            c = int(h[k:k + 2], 16) / 255
            v.append(c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4)
        return 0.2126 * v[0] + 0.7152 * v[1] + 0.0722 * v[2]
    a, b = lum(src_theme_fg), lum(src_theme_bg)
    hi, lo = max(a, b), min(a, b)
    return (hi + 0.05) / (lo + 0.05)


def main():
    if len(sys.argv) < 2:
        print('使い方: python verify.py deck.pptx [出力先]')
        return 2
    src = sys.argv[1]
    out = os.path.abspath(sys.argv[2] if len(sys.argv) > 2 else 'slides_png')
    if not os.path.exists(src):
        print('見つからない: %s' % src)
        return 2

    print('== 構造チェック ==')
    prs, issues, empty_notes = structural(src)
    total = len(prs.slides._sldIdLst)
    for x in issues:
        print('  ! ' + x)
    if empty_notes:
        print('  ! ノートが空: %s 枚目（台本はノートが正。空欄は書き忘れの可能性）'
              % ','.join(map(str, empty_notes)))
    if not issues and not empty_notes:
        print('  問題なし（%d枚）' % total)
    print('  ⚠ ただし構造チェックは**折り返しの崩れを捕まえない**。以下が本番。')

    print('== 描画して絵にする ==')
    _clean(out)
    n = by_powerpoint(src, out)
    engine = 'PowerPoint 本体'
    if n is None:
        n = by_libreoffice(src, out)
        engine = 'LibreOffice（書体差が出うる）'
    if n is None:
        print()
        print('  ✖ **目視検査していません。** PowerPoint も LibreOffice も使えませんでした。')
        print('     この状態で「完成」と言ってはいけません。次のどれかをしてください:')
        print('       - Office の入った PC でこのスクリプトを流す')
        print('       - LibreOffice ＋ pip install pymupdf を入れる')
        print('       - 人間が %s を開いて全%d枚を目で確認する' % (src, total))
        return 1

    print('  %s で %d枚を書き出した -> %s' % (engine, n, out))
    print()
    print('  ★ ここで終わりではない。**書き出した%d枚を1枚ずつ開いて目で見る。**' % n)
    print('    見るのは主にこの4点:')
    print('      1. 文の途中で折り返して、末尾の数文字だけが次行に残っていないか')
    print('      2. 文字が図形や別の文字に重なっていないか')
    print('      3. 同じ意味の要素（棒・列・行）の起点と長さが揃っているか')
    print('      4. 会場の後ろから読める大きさか（縮小表示して読めるかで判断）')
    return 0


if __name__ == '__main__':
    sys.exit(main())
